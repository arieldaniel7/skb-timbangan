"""Ekspor laporan: CSV, Excel (openpyxl), PDF/cetak (Qt). Semuanya lokal."""
import csv
import html

from app import util
from app.errors import UserError


def export_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:  # ';' agar Excel Indonesia membukanya benar
        w = csv.writer(fh, delimiter=";")
        w.writerow(headers)
        for r in rows:
            w.writerow(["" if v is None else v for v in r])


def export_xlsx(path, title, headers, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise UserError("Modul 'openpyxl' belum terpasang. Gunakan CSV atau pasang openpyxl.")
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Laporan")[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [10])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(40, width + 2)
    wb.save(path)


def table_html(title, subtitle, headers, rows, total_text=""):
    def cell(v):
        if v is None:
            return ""
        if isinstance(v, int):
            return util.format_int(v)
        if isinstance(v, float):
            return ("%.2f" % v).replace(".", ",")
        return html.escape(str(v))

    parts = ['<html><body style="font-family:Arial; font-size:9pt;"><h3>%s</h3><p>%s</p>' %
             (html.escape(title), html.escape(subtitle)),
             '<table width="100%" border="1" cellspacing="0" cellpadding="3"><tr>']
    parts += ['<th bgcolor="#dddddd">%s</th>' % html.escape(h) for h in headers]
    parts.append("</tr>")
    for r in rows:
        parts.append("<tr>" + "".join(
            '<td align="%s">%s</td>' % ("right" if isinstance(v, (int, float)) else "left", cell(v)) for v in r) + "</tr>")
    parts.append("</table>")
    if total_text:
        parts.append("<p><b>%s</b></p>" % html.escape(total_text))
    parts.append("</body></html>")
    return "".join(parts)


def export_pdf(path, html_text):
    from PyQt5.QtGui import QTextDocument
    from PyQt5.QtPrintSupport import QPrinter
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setOrientation(QPrinter.Landscape)
    doc = QTextDocument()
    doc.setHtml(html_text)
    doc.print_(printer)


def print_html(html_text, parent=None):
    from PyQt5.QtGui import QTextDocument
    from PyQt5.QtPrintSupport import QPrinter, QPrintDialog
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOrientation(QPrinter.Landscape)
    dlg = QPrintDialog(printer, parent)
    if dlg.exec_() != QPrintDialog.Accepted:
        return False
    doc = QTextDocument()
    doc.setHtml(html_text)
    doc.print_(printer)
    return True
